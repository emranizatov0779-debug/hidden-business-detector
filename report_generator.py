"""
report_generator.py  v2
========================
Генерация отчётов для регуляторов (PDF и Excel).

ИСПРАВЛЕНО v2:
  Добавлена diagnose_environment() — показывает РЕАЛЬНЫЙ путь Python,
  который использует запущенный Streamlit-процесс, и точную команду
  pip install ИМЕННО для этого интерпретатора.

  Главная причина "я ставил pip install reportlab, но не работает":
  на компьютере несколько Python (системный, venv, conda) — `pip install`
  в терминале мог поставить пакет в ДРУГОЙ Python, не в тот, из которого
  запущен `streamlit run app.py`.
"""

import os
import io
import sys
import json
import subprocess
from datetime import datetime


# ─── ДИАГНОСТИКА ОКРУЖЕНИЯ ────────────────────────────────────────────────────

def _is_externally_managed() -> bool:
    """
    Определяет PEP 668 (externally-managed-environment) — актуально для
    Homebrew Python на macOS. В этом случае голый `pip install` всегда
    падает с ошибкой, даже если путь к python указан верно.
    """
    base_dir = os.path.dirname(sys.executable)
    # marker файл лежит рядом с интерпретатором, в lib/pythonX.Y/EXTERNALLY-MANAGED
    candidates = [
        os.path.join(os.path.dirname(base_dir), 'lib',
                     f'python{sys.version_info.major}.{sys.version_info.minor}',
                     'EXTERNALLY-MANAGED'),
    ]
    return any(os.path.exists(c) for c in candidates)


def diagnose_environment() -> dict:
    """Проверяет библиотеки именно в том Python, из которого запущен Streamlit."""
    info = {
        'python_executable': sys.executable,
        'python_version': sys.version.split()[0],
        'externally_managed': _is_externally_managed(),
        'in_venv': sys.prefix != sys.base_prefix,
        'libraries': {},
    }
    for lib_name, import_name in [('reportlab', 'reportlab'), ('openpyxl', 'openpyxl')]:
        try:
            mod = __import__(import_name)
            info['libraries'][lib_name] = {
                'installed': True,
                'version': getattr(mod, '__version__', 'unknown'),
                'location': getattr(mod, '__file__', 'unknown'),
            }
        except ImportError as e:
            info['libraries'][lib_name] = {'installed': False, 'error': str(e)}
    return info


def get_install_command(lib_name: str) -> str:
    """
    Команда установки, адаптированная под окружение:
      - Homebrew/system Python (PEP 668)  → добавляет --break-system-packages
      - Обычный Python / venv             → обычный pip install
    """
    base = f'"{sys.executable}" -m pip install {lib_name}'
    if _is_externally_managed():
        return base + ' --break-system-packages --user'
    return base


def get_venv_setup_commands(libs: list[str]) -> str:
    """
    Альтернативный (рекомендуемый) путь: создать отдельный venv для проекта,
    чтобы больше никогда не упираться в externally-managed-environment.
    """
    libs_str = ' '.join(libs)
    return (
        f'cd "{os.getcwd()}"\n'
        f'python3 -m venv venv\n'
        f'source venv/bin/activate\n'
        f'pip install streamlit pandas plotly {libs_str}\n'
        f'streamlit run app.py'
    )


# ─── PDF ОТЧЁТ ───────────────────────────────────────────────────────────────

def generate_pdf_report(df_stats: dict, patterns: list, ext_intel: dict = None) -> bytes | None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.colors import HexColor, white
        from reportlab.lib.units import cm
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                         Table, TableStyle, HRFlowable, PageBreak)
    except ImportError:
        return None

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm,
                            title='Hidden Business Detector — Отчёт')

    styles = getSampleStyleSheet()
    BLUE, RED, ORANGE = HexColor('#1565c0'), HexColor('#b71c1c'), HexColor('#e65100')
    GREY, LIGHT = HexColor('#64748b'), HexColor('#e8f0fe')

    title_style = ParagraphStyle('Title', parent=styles['Title'], textColor=BLUE, fontSize=18, spaceAfter=6)
    h2_style    = ParagraphStyle('H2', parent=styles['Heading2'], textColor=BLUE, fontSize=13, spaceBefore=16, spaceAfter=6)
    h3_style    = ParagraphStyle('H3', parent=styles['Heading3'], textColor=HexColor('#0d3b8e'), fontSize=11, spaceBefore=10, spaceAfter=4)
    small_style = ParagraphStyle('Small', parent=styles['Normal'], fontSize=8, leading=11, textColor=GREY)
    risk_colors = {'CRITICAL': RED, 'HIGH': ORANGE, 'MEDIUM': HexColor('#f57f17')}

    story = []
    now = datetime.now().strftime('%d.%m.%Y %H:%M')

    story.append(Spacer(1, 1*cm))
    story.append(Paragraph('🔍 Hidden Business Detector', title_style))
    story.append(Paragraph('Отчёт об обнаруженных схемах скрытого бизнеса', h3_style))
    story.append(Paragraph(f'Дата формирования: {now}', small_style))
    story.append(Paragraph('Cashless Avengers · Confidential', small_style))
    story.append(HRFlowable(width='100%', thickness=2, color=BLUE, spaceAfter=16))

    story.append(Paragraph('Сводная статистика', h2_style))
    kpi_data = [
        ['Показатель', 'Значение'],
        ['Всего мерчантов проверено', f"{df_stats.get('total', 0):,}"],
        ['HOT LEAD (риск > 75%)', f"{df_stats.get('hot', 0):,}  ({df_stats.get('hot',0)/max(df_stats.get('total',1),1)*100:.1f}%)"],
        ['WARM LEAD (риск 50–75%)', f"{df_stats.get('warm', 0):,}  ({df_stats.get('warm',0)/max(df_stats.get('total',1),1)*100:.1f}%)"],
        ['LOW риск (< 50%)', f"{df_stats.get('low', 0):,}"],
        ['Средняя вероятность скрытого бизнеса', f"{df_stats.get('avg_prob', 0):.1f}%"],
        ['Обнаружено поведенческих схем', str(len(patterns))],
    ]
    kpi_table = Table(kpi_data, colWidths=[10*cm, 7*cm])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), BLUE), ('TEXTCOLOR', (0,0), (-1,0), white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'), ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [LIGHT, white]),
        ('GRID', (0,0), (-1,-1), 0.5, HexColor('#c5d8f0')), ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph('Обнаруженные схемы скрытого бизнеса', h2_style))
    critical_pats = [p for p in patterns if p.get('risk') == 'CRITICAL']
    high_pats     = [p for p in patterns if p.get('risk') == 'HIGH']
    medium_pats   = [p for p in patterns if p.get('risk') == 'MEDIUM']

    for risk_group, group_label in [(critical_pats, '🔴 CRITICAL'), (high_pats, '🟠 HIGH'), (medium_pats, '🟡 MEDIUM')]:
        if not risk_group:
            continue
        story.append(Paragraph(f'{group_label} ({len(risk_group)} схем)', h3_style))
        for p in risk_group:
            pat_data = [
                ['Схема:', p.get('name', '—')],
                ['Мерчантов:', f"{p.get('matched_count', 0):,} ({p.get('match_rate_pct', 0)}%)"],
                ['MCC:', p.get('dominant_mcc', 'N/A')],
                ['Описание:', p.get('description', '—')],
                ['Обнаружение:', p.get('prevention', '—')],
                ['Действие банка:', p.get('bank_action', '—')],
            ]
            pat_table = Table(pat_data, colWidths=[4*cm, 13*cm])
            r = p.get('risk', 'MEDIUM')
            pat_table.setStyle(TableStyle([
                ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'), ('FONTSIZE', (0,0), (-1,-1), 8.5),
                ('TEXTCOLOR', (0,0), (0,-1), risk_colors.get(r, BLUE)),
                ('BACKGROUND', (0,0), (-1,0), HexColor('#f8fafc')),
                ('GRID', (0,0), (-1,-1), 0.3, HexColor('#e2e8f0')), ('PADDING', (0,0), (-1,-1), 5),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ]))
            story.append(pat_table)
            story.append(Spacer(1, 0.3*cm))

    if ext_intel and ext_intel.get('patterns'):
        story.append(PageBreak())
        story.append(Paragraph('Глобальные угрозы (External Intelligence)', h2_style))
        story.append(Paragraph(
            f'Источник: {"Claude AI" if ext_intel.get("source")=="claude_api" else "База знаний"} · '
            f'Обновлено: {ext_intel.get("generated_at","")[:16]}', small_style))
        story.append(Spacer(1, 0.3*cm))
        for threat in ext_intel.get('threats', [])[:5]:
            r = threat.get('risk', 'MEDIUM')
            threat_data = [
                ['Угроза:', threat.get('name', '—')], ['Риск:', r],
                ['Регион:', threat.get('region', '—')],
                ['Описание:', threat.get('description', '—')],
                ['Источник:', threat.get('source', '—')],
            ]
            t = Table(threat_data, colWidths=[4*cm, 13*cm])
            t.setStyle(TableStyle([
                ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'), ('FONTSIZE', (0,0), (-1,-1), 8.5),
                ('TEXTCOLOR', (0,0), (0,-1), risk_colors.get(r, BLUE)),
                ('GRID', (0,0), (-1,-1), 0.3, HexColor('#e2e8f0')), ('PADDING', (0,0), (-1,-1), 5),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.25*cm))

    story.append(Spacer(1, 1*cm))
    story.append(HRFlowable(width='100%', thickness=1, color=HexColor('#c5d8f0')))
    story.append(Paragraph(f'Hidden Business Detector · Cashless Avengers 2026 · {now} · CONFIDENTIAL', small_style))

    doc.build(story)
    return buf.getvalue()


# ─── EXCEL ОТЧЁТ ─────────────────────────────────────────────────────────────

def generate_excel_report(df_merchants, patterns: list, df_stats: dict) -> bytes | None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    wb = openpyxl.Workbook()
    BLUE_FILL   = PatternFill('solid', fgColor='1565C0')
    RED_FILL    = PatternFill('solid', fgColor='B71C1C')
    ORANGE_FILL = PatternFill('solid', fgColor='E65100')
    YELLOW_FILL = PatternFill('solid', fgColor='F57F17')
    LIGHT_FILL  = PatternFill('solid', fgColor='E8F0FE')
    HEADER_FONT = Font(bold=True, color='FFFFFF', size=10)
    thin   = Side(style='thin', color='C5D8F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(cell, fill=BLUE_FILL):
        cell.font = HEADER_FONT; cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    def style_cell(cell, bold=False):
        cell.font = Font(bold=bold, size=9)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cell.border = border

    ws1 = wb.active
    ws1.title = '📊 Сводка'
    ws1.column_dimensions['A'].width = 40
    ws1.column_dimensions['B'].width = 20
    rows_summary = [
        ('Hidden Business Detector — Отчёт', None), ('', None),
        ('Показатель', 'Значение'),
        ('Дата формирования', datetime.now().strftime('%d.%m.%Y %H:%M')), ('', None),
        ('Всего мерчантов', df_stats.get('total', 0)),
        ('HOT LEAD (> 75%)', df_stats.get('hot', 0)),
        ('WARM LEAD (50–75%)', df_stats.get('warm', 0)),
        ('LOW риск (< 50%)', df_stats.get('low', 0)),
        ('Средняя вероятность', f"{df_stats.get('avg_prob', 0):.1f}%"), ('', None),
        ('Схем (CRITICAL)', sum(1 for p in patterns if p.get('risk')=='CRITICAL')),
        ('Схем (HIGH)', sum(1 for p in patterns if p.get('risk')=='HIGH')),
        ('Схем (MEDIUM)', sum(1 for p in patterns if p.get('risk')=='MEDIUM')),
    ]
    for i, (label, value) in enumerate(rows_summary, 1):
        ws1[f'A{i}'] = label
        if value is not None: ws1[f'B{i}'] = value
        if i == 1: ws1[f'A{i}'].font = Font(bold=True, size=14, color='1565C0')
        elif i == 3: style_header(ws1[f'A{i}']); style_header(ws1[f'B{i}'])
        elif label and value is not None:
            style_cell(ws1[f'A{i}'], bold=True); style_cell(ws1[f'B{i}'])

    ws2 = wb.create_sheet('🔍 Паттерны')
    headers2 = ['Риск','Название схемы','Тип','Мерчантов','Доля %','MCC','Описание','Как обнаружить','Действие банка']
    col_widths2 = [10,35,12,12,10,10,50,40,40]
    for j, (h, w) in enumerate(zip(headers2, col_widths2), 1):
        style_header(ws2.cell(row=1, column=j, value=h))
        ws2.column_dimensions[get_column_letter(j)].width = w
    for i, p in enumerate(sorted(patterns, key=lambda x: {'CRITICAL':0,'HIGH':1,'MEDIUM':2}.get(x.get('risk'),3)), 2):
        r = p.get('risk', 'MEDIUM')
        fill = {'CRITICAL':RED_FILL,'HIGH':ORANGE_FILL,'MEDIUM':YELLOW_FILL}.get(r, LIGHT_FILL)
        row_data = [r, p.get('name','—'),
                   'AI-кластер' if p.get('is_ai') else ('Новый' if p.get('type')=='discovered' else 'Известный'),
                   p.get('matched_count',0), p.get('match_rate_pct',0), p.get('dominant_mcc','N/A'),
                   p.get('description','—'), p.get('prevention','—'), p.get('bank_action','—')]
        for j, val in enumerate(row_data, 1):
            cell = ws2.cell(row=i, column=j, value=val)
            style_cell(cell)
            if j == 1: cell.fill = fill; cell.font = Font(bold=True, color='FFFFFF', size=9)
        ws2.row_dimensions[i].height = 45
    ws2.freeze_panes = 'A2'
    if patterns:
        ws2.auto_filter.ref = f'A1:I{len(patterns)+1}'

    if df_merchants is not None and len(df_merchants) > 0:
        ws3 = wb.create_sheet('📋 Мерчанты')
        show_cols = [c for c in ['merchant_id','mcc','dominant_mcc','country','dominant_country','probability','tier'] if c in df_merchants.columns]
        headers3 = {'merchant_id':'ID Мерчанта','mcc':'MCC','dominant_mcc':'MCC','country':'Страна','dominant_country':'Страна','probability':'Вер-ть %','tier':'Тир'}
        for j, col in enumerate(show_cols, 1):
            style_header(ws3.cell(row=1, column=j, value=headers3.get(col, col)))
            ws3.column_dimensions[get_column_letter(j)].width = 18
        hot_df = df_merchants[df_merchants.get('tier','')=='HOT'].head(500) if 'tier' in df_merchants.columns else df_merchants.head(500)
        if 'probability' in hot_df.columns:
            hot_df = hot_df.sort_values('probability', ascending=False)
        for i, (_, row) in enumerate(hot_df[show_cols].iterrows(), 2):
            tier = row.get('tier', '')
            row_fill = PatternFill('solid', fgColor='FEE2E2') if tier=='HOT' else \
                      PatternFill('solid', fgColor='FFF3E0') if tier=='Warm' else \
                      PatternFill('solid', fgColor='F0FDF4')
            for j, col in enumerate(show_cols, 1):
                cell = ws3.cell(row=i, column=j, value=row.get(col, ''))
                style_cell(cell); cell.fill = row_fill
        ws3.freeze_panes = 'A2'
        ws3.auto_filter.ref = f'A1:{get_column_letter(len(show_cols))}1'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
